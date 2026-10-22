#!/usr/bin/env python3
"""Create the corrected Review 1 PPTX for the switch-based PLL charge pump."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "exports" / "Switch_Based_PLL_Charge_Pump_Review_1.pptx"
CIRCUIT = ROOT / "exports" / "switch_based_charge_pump_circuit.png"

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(15, 28, 35)
INK = RGBColor(17, 29, 35)
MUTED = RGBColor(82, 104, 112)
LINE = RGBColor(201, 215, 216)
PALE = RGBColor(242, 247, 246)
PALE_BLUE = RGBColor(233, 242, 250)
PALE_GREEN = RGBColor(236, 246, 235)
PALE_AMBER = RGBColor(251, 244, 226)
CYAN = RGBColor(35, 136, 139)
BLUE = RGBColor(52, 108, 159)
GREEN = RGBColor(52, 101, 56)
AMBER = RGBColor(149, 100, 0)
RED = RGBColor(159, 47, 45)
WHITE = RGBColor(255, 255, 255)


def color(fill, rgb):
    fill.solid()
    fill.fore_color.rgb = rgb


def textbox(slide, x, y, w, h, text="", size=14, color_rgb=INK, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color_rgb
    return box


def add_lines(tf, lines, size=14, color_rgb=INK, bullet=False, gap=5):
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        p.text = item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color_rgb
        if bullet:
            p.text = "• " + item


def card(slide, x, y, w, h, title, body=None, accent=CYAN, fill=WHITE, title_size=14):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    color(shape.fill, fill)
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.8)
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    color(accent_bar.fill, accent)
    accent_bar.line.fill.background()
    textbox(slide, x + 0.23, y + 0.18, w - 0.42, 0.30, title, title_size, INK, True)
    if body:
        box = slide.shapes.add_textbox(Inches(x + 0.23), Inches(y + 0.60), Inches(w - 0.44), Inches(h - 0.72))
        add_lines(box.text_frame, body, size=11.5, color_rgb=MUTED, bullet=True, gap=5)
    return shape


def base_slide(prs, num, category, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    color(background.fill, RGBColor(250, 252, 251))
    textbox(slide, 0.55, 0.34, 8.0, 0.22, category.upper(), 8.5, MUTED, True, "Aptos Mono")
    textbox(slide, 0.55, 0.62, 11.9, 0.58, title, 25, INK, True)
    textbox(slide, 0.55, 1.22, 11.8, 0.32, subtitle, 11.5, MUTED)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.55), Inches(1.68), Inches(12.78), Inches(1.68))
    line.line.color.rgb = LINE
    line.line.width = Pt(0.8)
    textbox(slide, 11.76, 0.42, 0.72, 0.28, f"{num:02d} / 10", 9, MUTED, True, "Aptos Mono", PP_ALIGN.RIGHT)
    textbox(slide, 0.55, 7.08, 7.2, 0.18, "Switch-Based CMOS Charge Pump - Review 1", 8.5, MUTED, False, "Aptos Mono")
    textbox(slide, 11.45, 7.08, 1.32, 0.18, "Cadence Virtuoso", 8.5, MUTED, False, "Aptos Mono", PP_ALIGN.RIGHT)
    return slide


def stat(slide, x, y, w, label, value, fill=PALE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.92))
    color(shape.fill, fill)
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.7)
    textbox(slide, x + .16, y + .16, w - .32, .18, label.upper(), 7.5, MUTED, True, "Aptos Mono")
    textbox(slide, x + .16, y + .42, w - .32, .27, value, 13.5, INK, True)


def draw_pulse(slide, x, y, w, label, points, stroke):
    textbox(slide, x, y - .05, .55, .18, label, 9, MUTED, True, "Aptos Mono")
    base = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + .65), Inches(y + .25), Inches(x + w), Inches(y + .25))
    base.line.color.rgb = LINE
    base.line.width = Pt(0.75)
    scaled = [(x + .65 + p[0] * (w - .65), y + .25 - p[1] * .32) for p in points]
    for a, b in zip(scaled, scaled[1:]):
        seg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1]))
        seg.line.color.rgb = stroke
        seg.line.width = Pt(2.2)


def create():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    prs.core_properties.title = "Switch-Based PLL CMOS Charge Pump - Review 1"
    prs.core_properties.subject = "Cadence Virtuoso simulation plan"
    prs.core_properties.author = "Pragyan Jain, Shivam Kumar, K Navin Aakash"

    # 1
    slide = base_slide(prs, 1, "Project review 1", "Switch-Based CMOS Charge Pump for PLL Control", "UP/DN charge redistribution creates a controllable Vop within the supply range.")
    textbox(slide, .72, 1.98, 5.55, 1.2, "Vop does not boost above VDD.", 31, NAVY, True)
    textbox(slide, .75, 3.18, 5.25, .62, "The selected S1-S4 circuit adds or removes small packets of charge at the loop-filter output.", 15, MUTED)
    stat(slide, .75, 4.18, 1.55, "Control command", "UP / DN", PALE_BLUE)
    stat(slide, 2.48, 4.18, 1.55, "Output range", "0 ≤ Vop < VDD", PALE_GREEN)
    stat(slide, 4.21, 4.18, 1.55, "Project scope", "Cadence only", PALE_AMBER)
    card(slide, 6.65, 2.05, 5.88, 3.95, "Reference paper", [
        "M. Sivasakthi and P. Radhika",
        "High-Speed MCML Charge Pump Design at 10 GHz Frequency in 45 nm CMOS Technology for PLL Application",
        "Paper reference: 1 V supply, 0-980 mV control-voltage swing",
        "Those are reference-paper results, not results claimed for this project"
    ], BLUE, PALE_BLUE, 15)
    textbox(slide, .75, 5.48, 5.25, .42, "Team: Pragyan Jain (24BLC1100), Shivam Kumar (24BLC1103), K Navin Aakash (23BLC1238)", 10.5, MUTED)

    # 2
    slide = base_slide(prs, 2, "Theory and PLL context", "From UP/DN Decisions to a Control Voltage", "The phase-frequency detector requests correction. The charge pump translates that request into Vop.")
    card(slide, .65, 2.02, 3.75, 3.78, "Signal chain", ["PFD creates UP or DN pulses", "S1-S4 select the charge route", "Ci transfers a small charge packet", "CL stores and smooths Vop", "Vop controls the PLL loop path"], CYAN, WHITE)
    card(slide, 4.78, 2.02, 3.75, 3.78, "What UP, DN, and hold mean", ["UP: deliver charge to CL", "DN: remove charge from CL", "Hold: isolate Vop as much as possible", "Vop always remains below VDD"], GREEN, PALE_GREEN)
    card(slide, 8.91, 2.02, 3.75, 3.78, "Core relation", ["A capacitor stores charge", "Q = C V", "More transferred charge produces a larger Vop step", "CL trades fast response for smooth output"], AMBER, PALE_AMBER)

    # 3
    slide = base_slide(prs, 3, "VLSI implementation", "What Will Be Built in Cadence Virtuoso", "The project recreates the selected transistor-switch circuit and tests it with a reusable schematic symbol and testbench.")
    card(slide, .65, 2.04, 3.7, 3.9, "Circuit implementation", ["Individual NMOS and PMOS switches for S1-S4", "Internal capacitor Ci and output capacitor CL", "VDD, ground, UP, DN, and complementary controls", "Clearly labelled Vop measurement node"], BLUE, WHITE)
    card(slide, 4.82, 2.04, 3.7, 3.9, "Testbench", ["DC supply and synchronized pulse sources", "Complementary inputs and optional timing gap", "Adjustable output load", "Vop and supply-current probes"], CYAN, PALE_BLUE)
    card(slide, 8.99, 2.04, 3.7, 3.9, "Simulation-only scope", ["Transient analyses in Cadence Virtuoso", "No chip fabrication or laboratory testing", "No final waveform claim in Review 1", "Later reviews add measured Cadence evidence"], AMBER, PALE_AMBER)

    # 4
    slide = base_slide(prs, 4, "Selected topology", "The S1-S4 Switch-Based PLL Charge Pump", "Ci moves charge through controlled paths. CL stores the resulting Vop below VDD.")
    if CIRCUIT.exists():
        slide.shapes.add_picture(str(CIRCUIT), Inches(.72), Inches(1.95), height=Inches(4.85))
    card(slide, 6.85, 2.06, 5.7, 1.15, "VDD and ground", ["Energy source and lower reference for the switching network"], BLUE, PALE_BLUE)
    card(slide, 6.85, 3.48, 5.7, 1.15, "UP / DN", ["Commands that request controlled upward or downward Vop movement"], CYAN, WHITE)
    card(slide, 6.85, 4.90, 5.7, 1.15, "Ci / CL", ["Transfer capacitor and output storage capacitor, respectively"], GREEN, PALE_GREEN)

    # 5
    slide = base_slide(prs, 5, "Operating sequence", "UP, DN, and Hold Control Vop", "The waveform below is a conceptual explanation of the planned simulation, not a completed Cadence result.")
    card(slide, .68, 2.02, 3.1, 3.85, "UP", ["UP and its complement activate the upper transfer arrangement", "Charge is delivered toward CL", "Vop rises by a small amount but stays below VDD"], GREEN, PALE_GREEN)
    card(slide, 4.08, 2.02, 3.1, 3.85, "DN", ["DN and its complement activate the lower transfer arrangement", "Charge leaves CL through the controlled path", "Vop falls by a small amount"], RED, WHITE)
    card(slide, 7.48, 2.02, 3.1, 3.85, "Hold", ["Neither transfer is requested", "The output is isolated as much as possible", "CL holds charge, while leakage can cause slow drift"], BLUE, PALE_BLUE)
    wave = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.88), Inches(2.02), Inches(1.77), Inches(3.85))
    color(wave.fill, PALE)
    wave.line.color.rgb = LINE
    textbox(slide, 11.05, 2.26, 1.4, .2, "CONCEPTUAL", 8, MUTED, True, "Aptos Mono", PP_ALIGN.CENTER)
    draw_pulse(slide, 11.02, 3.03, 1.42, "UP", [(0, 0), (.20, 0), (.20, 1), (.50, 1), (.50, 0), (1, 0)], GREEN)
    draw_pulse(slide, 11.02, 3.88, 1.42, "DN", [(0, 0), (.55, 0), (.55, 1), (.78, 1), (.78, 0), (1, 0)], RED)
    draw_pulse(slide, 11.02, 4.73, 1.42, "Vop", [(0, 0), (.18, 0), (.32, .38), (.52, .38), (.70, .15), (1, .15)], BLUE)
    textbox(slide, 11.02, 5.45, 1.42, .28, "0 ≤ Vop < VDD", 8, BLUE, True, "Aptos Mono", PP_ALIGN.CENTER)

    # 6
    slide = base_slide(prs, 6, "Design variables", "What We Will Change, One Item at a Time", "A documented base case makes every comparison explainable.")
    variables = [
        ("S1-S4 size", "Response speed, power, and switching disturbance", BLUE),
        ("Ci", "Charge moved per UP or DN action", CYAN),
        ("CL", "Ripple, hold behavior, and settling time", GREEN),
        ("UP/DN timing", "Pulse width, rate, safe gap, and overlap", AMBER),
        ("Output load", "Step size, drift, and stability", RED),
        ("VDD, temperature, model", "Robustness across PVT conditions", BLUE),
    ]
    for i, (name, body, accent) in enumerate(variables):
        col, row = i % 3, i // 3
        card(slide, .68 + col * 4.12, 2.05 + row * 2.05, 3.78, 1.65, name, [body], accent, WHITE)

    # 7
    slide = base_slide(prs, 7, "Cadence analysis plan", "Measurements That Define Useful PLL Charge-Pump Behavior", "The final choice will balance response, smoothness, hold performance, and power use.")
    measures = [
        ("UP / DN step", "Vop movement for a command"), ("Rise / fall / settling", "Speed of the requested response"),
        ("Ripple / spikes", "Output smoothness and switching disturbance"), ("Hold drift", "Vop change with no active command"),
        ("Supply current", "Current from VDD during each state"), ("Energy comparison", "Fair comparison under identical test conditions"),
    ]
    for i, (name, body) in enumerate(measures):
        col, row = i % 3, i // 3
        card(slide, .68 + col * 4.12, 2.05 + row * 1.85, 3.78, 1.42, name, [body], [BLUE, CYAN, GREEN, AMBER, RED, BLUE][i], [PALE_BLUE, WHITE, PALE_GREEN, PALE_AMBER, WHITE, PALE_BLUE][i])
    textbox(slide, .8, 5.92, 11.5, .35, "Fair comparison rule: keep the same starting Vop, load, input sequence, simulation length, and measurement points.", 12, INK, True, align=PP_ALIGN.CENTER)

    # 8
    slide = base_slide(prs, 8, "Robustness plan", "PVT and Repeated-Run Checks", "After individual trade-offs are understood, promising settings will be tested under changing conditions.")
    card(slide, .68, 2.08, 3.75, 3.9, "Process model", ["Typical, faster, and slower models from the selected Cadence library", "Observe correct UP/DN direction, step size, ripple, drift, and power"], BLUE, PALE_BLUE)
    card(slide, 4.79, 2.08, 3.75, 3.9, "Voltage and temperature", ["Lower, normal, and higher VDD", "Cold, room, and hot settings", "Check that Vop remains controlled within the supply range"], AMBER, PALE_AMBER)
    card(slide, 8.90, 2.08, 3.75, 3.9, "Combined and statistical checks", ["Example: low VDD with slow devices", "Example: hot temperature with heavier load", "Repeated runs only if supported by the library"], GREEN, PALE_GREEN)

    # 9
    slide = base_slide(prs, 9, "Simulation method", "From a Base Waveform to a Final Simulated Setup", "The base transient makes UP, DN, and hold visible before any parameter sweep begins.")
    steps = [("1", "Quiet", "Start from a known Vop condition."), ("2", "UP", "Apply a group of UP pulses."), ("3", "Hold", "Observe ripple and drift after UP."), ("4", "DN", "Apply a group of DN pulses."), ("5", "Hold", "Observe final stability and supply current.")]
    for i, (n, title, body) in enumerate(steps):
        x = .68 + i * 2.48
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.32), Inches(2.08), Inches(2.85))
        color(shape.fill, [PALE_BLUE, PALE_GREEN, PALE, PALE_AMBER, PALE][i])
        shape.line.color.rgb = LINE
        textbox(slide, x + .18, 2.53, .38, .28, n, 18, CYAN, True, "Aptos Mono")
        textbox(slide, x + .18, 3.02, 1.68, .3, title, 15, INK, True)
        textbox(slide, x + .18, 3.52, 1.65, .9, body, 10.5, MUTED)
        if i < 4:
            arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 2.1), Inches(3.75), Inches(x + 2.38), Inches(3.75))
            arrow.line.color.rgb = CYAN
            arrow.line.width = Pt(1.6)
    textbox(slide, .9, 5.8, 11.5, .42, "Save this waveform as the base case. Then vary one main parameter per simple comparison.", 12.5, INK, True, align=PP_ALIGN.CENTER)

    # 10
    slide = base_slide(prs, 10, "Project roadmap", "Three Reviews: Plan, Evidence, and Final Comparison", "Review 1 does not claim completed simulation results. Results are introduced only after Cadence runs are completed and checked.")
    card(slide, .68, 2.05, 3.75, 3.9, "Review 1 - Current", ["Explain the selected S1-S4 charge pump", "Define the Cadence-only scope", "Prepare the test procedure and analysis list", "Deliverable: approved design and simulation plan"], CYAN, PALE_BLUE)
    card(slide, 4.79, 2.05, 3.75, 3.9, "Review 2 - Implementation", ["Build the schematic, symbol, and testbench", "Run the first real UP, DN, and hold waveforms", "Begin switch, capacitor, timing, and load comparisons", "Deliverable: initial simulation results"], AMBER, PALE_AMBER)
    card(slide, 8.90, 2.05, 3.75, 3.9, "Review 3 - Conclusion", ["Complete PVT and combined-condition checks", "Compare useful cases and select the final simulated setup", "Present waveforms, measured trade-offs, and limitations", "Deliverable: final Cadence simulation report"], GREEN, PALE_GREEN)
    textbox(slide, .85, 6.30, 11.7, .36, "Reference paper: Sivasakthi and Radhika, A High-Speed MCML Charge Pump Design at 10 GHz Frequency in 45 nm CMOS Technology for PLL Application.", 9.5, MUTED, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    create()
