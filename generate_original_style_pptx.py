#!/usr/bin/env python3
"""Package the corrected high-resolution slide renders into a PPTX."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
IMAGES = ROOT / "exports" / "corrected_slides_png"
OUTPUT = ROOT / "exports" / "CMOS_Charge_Pump_Review_1_Corrected.pptx"


def main():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "CMOS Charge Pump Review 1 - Corrected"
    presentation.core_properties.subject = "Switch-based PLL charge pump, Vop below VDD"
    presentation.core_properties.author = "Pragyan Jain, Shivam Kumar, K Navin Aakash"

    for number in range(1, 11):
        suffix = f"0{number}" if number < 10 else "010"
        image = IMAGES / f"slide-{suffix}.png"
        if not image.exists():
            raise FileNotFoundError(image)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)

    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
